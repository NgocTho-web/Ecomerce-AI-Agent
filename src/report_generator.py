"""Tạo báo cáo phân tích Ecommerce-AI-Agent ở định dạng Word và PDF.

Chạy từ thư mục gốc dự án:
    python -m src.report_generator
    python -m src.report_generator --output-dir reports

Báo cáo dùng dữ liệu CLEANED_DATA_PATH trong src.config; AI nhận định lấy từ
src.agent.run_automatic_analysis. Word/PDF và biểu đồ được lưu trong thư mục output.
"""
from __future__ import annotations

import argparse
import logging
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd

LOG = logging.getLogger("report_generator")


def _project_root() -> Path:
    """Tìm thư mục gốc theo vị trí file src/report_generator.py."""
    return Path(__file__).resolve().parent.parent


def _load_data() -> tuple[pd.DataFrame, Path]:
    """Nạp dữ liệu từ cấu hình dự án, kiểm tra các cột tối thiểu."""
    try:
        from src.config import CLEANED_DATA_PATH
        data_path = Path(os.fspath(CLEANED_DATA_PATH))
    except Exception as exc:
        raise RuntimeError(
            "Không đọc được CLEANED_DATA_PATH từ src.config. "
            "Hãy kiểm tra file src/config.py."
        ) from exc

    if not data_path.is_absolute():
        data_path = (_project_root() / data_path).resolve()
    if not data_path.is_file():
        raise FileNotFoundError(
            f"Không tìm thấy dữ liệu sạch: {data_path}\n"
            "Hãy kiểm tra CLEANED_DATA_PATH trong src/config.py và bảo đảm "
            "file dữ liệu đã được tạo."
        )
    try:
        df = pd.read_csv(data_path)
    except Exception as exc:
        raise RuntimeError(f"Không thể đọc CSV '{data_path}': {exc}") from exc

    required = {"Sales", "Profit", "Quantity"}
    missing = sorted(required - set(df.columns))
    if missing:
        raise ValueError(
            f"Dữ liệu thiếu cột bắt buộc: {missing}. "
            f"Cột hiện có: {list(df.columns)}"
        )
    for col in ("Sales", "Profit", "Quantity", "Discount"):
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    if "Order Date" in df.columns:
        df["Order Date"] = pd.to_datetime(df["Order Date"], errors="coerce")
        df["YearMonth"] = df["Order Date"].dt.to_period("M").astype("string")
    return df, data_path


def _number(value: Any, currency: bool = False) -> str:
    try:
        if pd.isna(value):
            return "N/A"
        text = f"{float(value):,.2f}"
        return f"${text}" if currency else text
    except (TypeError, ValueError):
        return "N/A"


def _safe_name(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", text).strip("_") or "report"


def _make_charts(df: pd.DataFrame, chart_dir: Path) -> list[dict[str, str]]:
    """Tạo các biểu đồ PNG dựa trên toàn bộ dữ liệu hợp lệ."""
    chart_dir.mkdir(parents=True, exist_ok=True)
    charts: list[dict[str, str]] = []

    def save_bar(group_col: str, metric: str, title: str, filename: str) -> None:
        if group_col not in df.columns or metric not in df.columns:
            return
        temp = df[[group_col, metric]].dropna().copy()
        if temp.empty:
            return
        grouped = temp.groupby(group_col, dropna=False)[metric].sum().sort_values(ascending=False).head(15)
        if grouped.empty:
            return
        fig, ax = plt.subplots(figsize=(10, 5.5))
        try:
            grouped.sort_values().plot(kind="barh", ax=ax)
            ax.set_title(title)
            ax.set_xlabel(f"Tổng {metric}")
            ax.set_ylabel(group_col)
            fig.tight_layout()
            path = chart_dir / filename
            fig.savefig(path, dpi=160, bbox_inches="tight")
            charts.append({"title": title, "path": str(path)})
        finally:
            plt.close(fig)

    save_bar("Category", "Sales", "Doanh số theo danh mục", "sales_by_category.png")
    save_bar("Category", "Profit", "Lợi nhuận theo danh mục", "profit_by_category.png")

    if "YearMonth" in df.columns and "Sales" in df.columns:
        temp = df[["YearMonth", "Sales"]].dropna()
        if not temp.empty:
            grouped = temp.groupby("YearMonth", dropna=False)["Sales"].sum().sort_index()
            fig, ax = plt.subplots(figsize=(11, 5.5))
            try:
                ax.plot(grouped.index.astype(str), grouped.values, marker="o")
                ax.set_title("Xu hướng doanh số theo tháng")
                ax.set_xlabel("Năm-tháng")
                ax.set_ylabel("Tổng Sales ($)")
                ax.tick_params(axis="x", rotation=55)
                fig.tight_layout()
                path = chart_dir / "sales_by_month.png"
                fig.savefig(path, dpi=160, bbox_inches="tight")
                charts.append({"title": "Xu hướng doanh số theo tháng", "path": str(path)})
            finally:
                plt.close(fig)
    return charts


def _build_kpis(df: pd.DataFrame) -> list[tuple[str, str, str]]:
    kpis: list[tuple[str, str, str]] = []
    kpis.append(("Tổng doanh số (Sales)", _number(df["Sales"].sum()), "USD, theo dữ liệu nguồn"))
    kpis.append(("Tổng lợi nhuận (Profit)", _number(df["Profit"].sum()), "USD, theo dữ liệu nguồn"))
    kpis.append(("Tổng số lượng (Quantity)", _number(df["Quantity"].sum()), "Đơn vị sản phẩm"))
    if "Order ID" in df.columns:
        kpis.append(("Số đơn hàng", f"{df['Order ID'].nunique():,}", "Đếm Order ID duy nhất"))
    kpis.append(("Số dòng dữ liệu", f"{len(df):,}", "Số bản ghi trong CSV"))
    if df["Sales"].sum(skipna=True) != 0:
        margin = df["Profit"].sum(skipna=True) / df["Sales"].sum(skipna=True) * 100
        kpis.append(("Biên lợi nhuận", f"{margin:.2f}%", "Tổng Profit / tổng Sales"))
    if "Discount" in df.columns:
        kpis.append(("Giảm giá trung bình", f"{df['Discount'].mean():.2f}", "Trung bình cột Discount"))
    return kpis


def _basic_insights(df: pd.DataFrame) -> list[str]:
    insights: list[str] = []
    sales = df["Sales"].sum(skipna=True)
    profit = df["Profit"].sum(skipna=True)
    quantity = df["Quantity"].sum(skipna=True)
    insights.append(f"Tổng doanh số là {_number(sales)}; tổng lợi nhuận là {_number(profit)}; tổng số lượng bán là {_number(quantity)}.")
    if sales:
        insights.append(f"Biên lợi nhuận tổng hợp đạt {profit / sales * 100:.2f}% (tổng Profit chia tổng Sales). Đây là tỷ lệ tổng hợp, không phải trung bình biên lợi nhuận từng đơn hàng.")
    if "Category" in df.columns:
        grouped = df.groupby("Category", dropna=False)[["Sales", "Profit"]].sum(min_count=1)
        if not grouped.empty:
            best_sales = grouped["Sales"].idxmax()
            worst_profit = grouped["Profit"].idxmin()
            insights.append(f"Danh mục có doanh số cao nhất: {best_sales} ({_number(grouped.loc[best_sales, 'Sales'])}).")
            if grouped.loc[worst_profit, "Profit"] < 0:
                insights.append(f"Danh mục có tổng lợi nhuận thấp nhất là {worst_profit} ({_number(grouped.loc[worst_profit, 'Profit'])}), đang ghi nhận lợi nhuận âm trong dữ liệu.")
            else:
                insights.append(f"Danh mục có tổng lợi nhuận thấp nhất là {worst_profit} ({_number(grouped.loc[worst_profit, 'Profit'])}); giá trị này vẫn không âm.")
    if profit < 0:
        insights.append("Tổng lợi nhuận toàn bộ dữ liệu đang âm; cần kiểm tra cơ cấu chi phí, giảm giá và các nhóm sản phẩm gây lỗ.")
    elif profit > 0:
        insights.append("Tổng lợi nhuận toàn bộ dữ liệu đang dương; nên tiếp tục so sánh lợi nhuận giữa danh mục, khu vực và thời gian để tìm nguyên nhân.")
    if df["Sales"].isna().any() or df["Profit"].isna().any() or df["Quantity"].isna().any():
        insights.append("Một số bản ghi thiếu giá trị số ở Sales, Profit hoặc Quantity; các phép tổng hợp bỏ qua giá trị thiếu theo mặc định của pandas.")
    return insights


def _run_ai_analysis(model: str | None = None) -> tuple[list[dict[str, Any]], list[str], str | None]:
    """Gọi agent tương thích với API run_automatic_analysis trong src/agent.py."""
    try:
        from src.agent import run_automatic_analysis
        result = run_automatic_analysis(model=model) if model else run_automatic_analysis()
        reports = result.get("reports", []) if isinstance(result, dict) else []
        insights = []
        for report in reports:
            task = str(report.get("task", "Phân tích"))
            answer = str(report.get("answer", "")).strip()
            if answer:
                insights.append(f"{task}:\n{answer}")
        if not reports:
            return [], [], "Agent không trả về nhiệm vụ phân tích nào."
        return reports, insights, None
    except Exception as exc:
        return [], [], f"Không chạy được AI Agent ({type(exc).__name__}: {exc}). Báo cáo vẫn được tạo bằng thống kê trực tiếp từ dữ liệu."


def _add_word_table(doc: Any, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = str(header)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)


def _create_word(path: Path, df: pd.DataFrame, data_path: Path, kpis: list[tuple[str, str, str]],
                 basic_insights: list[str], ai_insights: list[str], ai_error: str | None,
                 charts: list[dict[str, str]], agent_reports: list[dict[str, Any]]) -> None:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise ImportError("Thiếu python-docx. Cài bằng: pip install python-docx") from exc

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("BÁO CÁO PHÂN TÍCH KINH DOANH E-COMMERCE", 0)
    doc.add_paragraph(f"Thời điểm tạo báo cáo: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    doc.add_paragraph(f"Nguồn dữ liệu: {data_path}")
    doc.add_paragraph(f"Phạm vi: toàn bộ {len(df):,} bản ghi trong tệp dữ liệu sạch.")

    doc.add_heading("1. Tổng quan và KPI", level=1)
    _add_word_table(doc, ["Chỉ số", "Giá trị", "Diễn giải"], [[a, b, c] for a, b, c in kpis])

    doc.add_heading("2. Nhận định từ dữ liệu", level=1)
    for item in basic_insights:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("3. Phân tích và nhận định AI Agent", level=1)
    if ai_error:
        doc.add_paragraph(ai_error)
    if ai_insights:
        for item in ai_insights:
            parts = item.split("\n", 1)
            doc.add_heading(parts[0], level=2)
            if len(parts) > 1:
                doc.add_paragraph(parts[1])
    else:
        doc.add_paragraph("Không có nhận định AI; tham khảo phần thống kê trực tiếp từ dữ liệu.")

    doc.add_heading("4. Biểu đồ", level=1)
    for chart in charts:
        chart_path = Path(chart["path"])
        if chart_path.is_file():
            doc.add_heading(chart["title"], level=2)
            doc.add_picture(str(chart_path), width=Inches(6.2))
        else:
            doc.add_paragraph(f"Không tìm thấy ảnh biểu đồ: {chart_path}")

    doc.add_heading("5. Trạng thái các nhiệm vụ AI", level=1)
    if agent_reports:
        _add_word_table(doc, ["Nhiệm vụ", "Trạng thái"], [[str(r.get("task", "")), "Thành công" if r.get("success") else "Có lỗi/Chưa hoàn tất"] for r in agent_reports])
    else:
        doc.add_paragraph("Không có kết quả nhiệm vụ AI.")
    doc.add_paragraph("Lưu ý: nhận định AI được tổng hợp từ kết quả công cụ; cần đối chiếu số liệu nguồn trước khi dùng cho quyết định kinh doanh.")
    doc.save(path)


def _register_pdf_font() -> tuple[str, str]:
    """Chọn font có hỗ trợ tiếng Việt nếu có sẵn trên máy."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates = [
        (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
    ]
    for regular, bold in candidates:
        if os.path.isfile(regular):
            pdfmetrics.registerFont(TTFont("ReportVietnamese", regular))
            bold_name = "ReportVietnamese"
            if os.path.isfile(bold):
                pdfmetrics.registerFont(TTFont("ReportVietnameseBold", bold))
                bold_name = "ReportVietnameseBold"
            return "ReportVietnamese", bold_name
    LOG.warning("Không tìm thấy font Unicode phổ biến; PDF có thể không hiển thị đủ dấu tiếng Việt.")
    return "Helvetica", "Helvetica-Bold"


def _create_pdf(path: Path, df: pd.DataFrame, data_path: Path, kpis: list[tuple[str, str, str]],
                basic_insights: list[str], ai_insights: list[str], ai_error: str | None,
                charts: list[dict[str, str]], agent_reports: list[dict[str, Any]]) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
    except ImportError as exc:
        raise ImportError("Thiếu reportlab. Cài bằng: pip install reportlab") from exc

    regular_font, bold_font = _register_pdf_font()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], fontName=bold_font, alignment=TA_CENTER, fontSize=17, leading=22, spaceAfter=12))
    styles.add(ParagraphStyle(name="VNHeading", parent=styles["Heading1"], fontName=bold_font, fontSize=13, leading=17, spaceBefore=10, spaceAfter=6))
    styles.add(ParagraphStyle(name="VNSubHeading", parent=styles["Heading2"], fontName=bold_font, fontSize=10.5, leading=14, spaceBefore=7, spaceAfter=4))
    styles.add(ParagraphStyle(name="VNBody", parent=styles["BodyText"], fontName=regular_font, fontSize=9, leading=13, spaceAfter=5, wordWrap="CJK"))
    styles.add(ParagraphStyle(name="VNSmall", parent=styles["BodyText"], fontName=regular_font, fontSize=7.5, leading=10, wordWrap="CJK"))

    def para(text: Any, style: str = "VNBody") -> Any:
        # Escape HTML special characters before creating Paragraph markup.
        from xml.sax.saxutils import escape
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), styles[style])

    story: list[Any] = [para("BÁO CÁO PHÂN TÍCH KINH DOANH E-COMMERCE", "ReportTitle")]
    story += [para(f"Thời điểm tạo báo cáo: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}") ,
              para(f"Nguồn dữ liệu: {data_path}"),
              para(f"Phạm vi: toàn bộ {len(df):,} bản ghi trong tệp dữ liệu sạch."), Spacer(1, 0.25 * cm),
              para("1. Tổng quan và KPI", "VNHeading")]
    kpi_data = [[para("Chỉ số", "VNSmall"), para("Giá trị", "VNSmall"), para("Diễn giải", "VNSmall")]]
    kpi_data += [[para(a, "VNSmall"), para(b, "VNSmall"), para(c, "VNSmall")] for a, b, c in kpis]
    table = Table(kpi_data, colWidths=[5.0 * cm, 3.5 * cm, 8.0 * cm], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")), ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story += [table, para("2. Nhận định từ dữ liệu", "VNHeading")]
    story.extend(para("• " + x) for x in basic_insights)
    story.append(para("3. Phân tích và nhận định AI Agent", "VNHeading"))
    if ai_error:
        story.append(para(ai_error))
    if ai_insights:
        for item in ai_insights:
            parts = item.split("\n", 1)
            story.append(para(parts[0], "VNSubHeading"))
            if len(parts) > 1:
                story.append(para(parts[1]))
    else:
        story.append(para("Không có nhận định AI; tham khảo phần thống kê trực tiếp từ dữ liệu."))
    story.append(para("4. Biểu đồ", "VNHeading"))
    for chart in charts:
        chart_path = Path(chart["path"])
        if chart_path.is_file():
            story.append(para(chart["title"], "VNSubHeading"))
            img = Image(str(chart_path))
            max_w, max_h = 17 * cm, 10 * cm
            scale = min(max_w / img.imageWidth, max_h / img.imageHeight, 1)
            img.drawWidth, img.drawHeight = img.imageWidth * scale, img.imageHeight * scale
            story += [img, Spacer(1, 0.2 * cm)]
    story.append(para("5. Trạng thái các nhiệm vụ AI", "VNHeading"))
    if agent_reports:
        rows = [[para("Nhiệm vụ", "VNSmall"), para("Trạng thái", "VNSmall")]]
        rows += [[para(r.get("task", ""), "VNSmall"), para("Thành công" if r.get("success") else "Có lỗi/Chưa hoàn tất", "VNSmall")] for r in agent_reports]
        t = Table(rows, colWidths=[11 * cm, 5.5 * cm], repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")), ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5)]))
        story.append(t)
    else:
        story.append(para("Không có kết quả nhiệm vụ AI."))
    story.append(Spacer(1, 0.3 * cm))
    story.append(para("Lưu ý: nhận định AI được tổng hợp từ kết quả công cụ; cần đối chiếu số liệu nguồn trước khi dùng cho quyết định kinh doanh."))

    def footer(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont(regular_font, 8)
        canvas.drawString(2 * cm, 1 * cm, "Ecommerce-AI-Agent | Báo cáo tự động")
        canvas.drawRightString(A4[0] - 2 * cm, 1 * cm, f"Trang {document.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=2 * cm, leftMargin=2 * cm, topMargin=1.7 * cm, bottomMargin=1.7 * cm, title="Báo cáo phân tích Ecommerce-AI-Agent", author="Ecommerce-AI-Agent")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def generate_report(output_dir: str | os.PathLike[str] | None = None, model: str | None = None) -> dict[str, Any]:
    """Tự động phân tích toàn bộ dữ liệu và xuất cả Word lẫn PDF."""
    root = _project_root()
    out_dir = Path(output_dir) if output_dir else root / "reports"
    if not out_dir.is_absolute():
        out_dir = (root / out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    chart_dir = out_dir / "charts"
    chart_dir.mkdir(parents=True, exist_ok=True)

    df, data_path = _load_data()
    if df.empty:
        raise ValueError(f"Tệp dữ liệu không có bản ghi: {data_path}")
    kpis = _build_kpis(df)
    basic_insights = _basic_insights(df)
    charts = _make_charts(df, chart_dir)
    agent_reports, ai_insights, ai_error = _run_ai_analysis(model=model)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    word_path = out_dir / f"ecommerce_report_{stamp}.docx"
    pdf_path = out_dir / f"ecommerce_report_{stamp}.pdf"
    _create_word(word_path, df, data_path, kpis, basic_insights, ai_insights, ai_error, charts, agent_reports)
    _create_pdf(pdf_path, df, data_path, kpis, basic_insights, ai_insights, ai_error, charts, agent_reports)

    return {
        "success": True,
        "rows_analyzed": int(len(df)),
        "source": str(data_path),
        "output_dir": str(out_dir),
        "word_path": str(word_path),
        "pdf_path": str(pdf_path),
        "charts": charts,
        "ai_tasks_completed": sum(1 for r in agent_reports if r.get("success")),
        "ai_tasks_total": len(agent_reports),
        "ai_warning": ai_error,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Tự động phân tích dữ liệu và xuất báo cáo Word/PDF.")
    parser.add_argument("--output-dir", default=None, help="Thư mục đầu ra (mặc định: reports trong thư mục gốc dự án).")
    parser.add_argument("--model", default=None, help="Tên model Ollama, ví dụ qwen2.5:7b.")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    try:
        result = generate_report(output_dir=args.output_dir, model=args.model)
    except Exception as exc:
        LOG.error("Không tạo được báo cáo: %s: %s", type(exc).__name__, exc)
        LOG.error("Kiểm tra dữ liệu, src.config, thư viện phụ thuộc và quyền ghi thư mục đầu ra.")
        return 1
    print("\nĐÃ TẠO BÁO CÁO THÀNH CÔNG")
    print(f"Số dòng phân tích : {result['rows_analyzed']:,}")
    print(f"Nguồn dữ liệu     : {result['source']}")
    print(f"Word              : {result['word_path']}")
    print(f"PDF               : {result['pdf_path']}")
    print(f"Biểu đồ PNG       : {len(result['charts'])}")
    print(f"Nhiệm vụ AI       : {result['ai_tasks_completed']}/{result['ai_tasks_total']}")
    if result.get("ai_warning"):
        print(f"Cảnh báo AI       : {result['ai_warning']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
