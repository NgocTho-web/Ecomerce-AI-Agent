"""
Chạy bộ kiểm thử và đánh giá hiệu năng AI Agent (Thành viên 4).

Yêu cầu: Ollama đang chạy và đã tải model (xem src/agent.py), đã chạy
    python tests/build_test_cases.py

Ví dụ:
    python tests/evaluate_agent.py                      # chạy tất cả, 3 lần/câu
    python tests/evaluate_agent.py --runs 1 --groups "Thống kê,Biểu đồ"
    python tests/evaluate_agent.py --ids TC01,TC05 --runs 5
    python tests/evaluate_agent.py --model qwen2.5:7b

Chỉ số xuất ra (tests/results/eval_<thời gian>.xlsx + .json):
    - Độ chính xác (Accuracy)  : % lượt chạy đạt điều kiện kiểm tra của test case
    - Chọn đúng công cụ        : % lượt chạy có gọi đúng tool kỳ vọng
    - Tính nhất quán           : % lượt chạy cho kết quả (Đúng/Sai) trùng với kết quả đa số của câu đó
    - Thời gian phản hồi       : trung bình, trung vị, P95, lớn nhất (giây)
    - Tỷ lệ lỗi                : lỗi hệ thống (không gọi được LLM/exception) và lỗi công cụ
"""
from __future__ import annotations

import argparse
import json
import math
import os
import re
import statistics
import sys
import time
import unicodedata
from datetime import datetime
from typing import Any, Callable, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

DEFAULT_CASES = os.path.join(BASE_DIR, "tests", "test_cases.xlsx")
RESULT_DIR = os.path.join(BASE_DIR, "tests", "results")

REFUSE_MARKERS = [
    "khong co du lieu", "khong tim thay", "khong ton tai", "khong co thong tin", "khong co ban ghi",
    "ngoai pham vi", "khong nam trong", "khong ho tro", "khong the", "khong duoc phep", "khong duoc",
    "xin loi", "rat tiec", "vui long", "lam ro", "cung cap them", "khong phu hop", "khong lien quan",
    "chi co du lieu", "khong co cong cu",
]
SYSTEM_FAILURE_MARKERS = ("khong goi duoc mo hinh ollama", "khong khoi tao duoc ollama")


# ----------------------------------------------------------------------------
# Tiện ích văn bản / số
# ----------------------------------------------------------------------------
def fold(text: Any) -> str:
    """Hạ chữ thường, bỏ dấu tiếng Việt để so khớp từ khóa."""
    s = unicodedata.normalize("NFD", str(text or "").lower().replace("đ", "d"))
    return "".join(ch for ch in s if unicodedata.category(ch) != "Mn")


_NUM_RE = re.compile(r"[-−–(]?\s?\$?\s?\d[\d.,]*(?:\s?(?:triệu|trieu|million|tỷ|ty|billion|nghìn|nghin|ngàn|ngan|k)\b)?",
                     re.IGNORECASE)
_SCALE = {"triệu": 1e6, "trieu": 1e6, "million": 1e6, "tỷ": 1e9, "ty": 1e9, "billion": 1e9,
          "nghìn": 1e3, "nghin": 1e3, "ngàn": 1e3, "ngan": 1e3, "k": 1e3}


def _parse_token(tok: str) -> list[tuple[float, int]]:
    """Trả về các cách hiểu (giá trị, số chữ số thập phân) của một chuỗi số."""
    out: list[tuple[float, int]] = []
    if "," in tok and "." in tok:
        if tok.rfind(",") > tok.rfind("."):           # 1.234,56
            out.append((float(tok.replace(".", "").replace(",", ".")), len(tok) - tok.rfind(",") - 1))
        else:                                          # 1,234.56
            out.append((float(tok.replace(",", "")), len(tok) - tok.rfind(".") - 1))
    elif "," in tok:
        if re.fullmatch(r"\d{1,3}(,\d{3})+", tok):     # 1,234,567
            out.append((float(tok.replace(",", "")), 0))
        else:                                          # 12,5
            out.append((float(tok.replace(",", ".")), len(tok) - tok.rfind(",") - 1))
    elif "." in tok:
        if re.fullmatch(r"\d{1,3}(\.\d{3})+", tok):    # 1.234.567 hoặc 2.297 (mơ hồ)
            out.append((float(tok.replace(".", "")), 0))
            if tok.count(".") == 1:
                out.append((float(tok), 3))
        else:
            out.append((float(tok), len(tok) - tok.rfind(".") - 1))
    else:
        out.append((float(tok), 0))
    return out


def extract_numbers(text: str) -> list[tuple[float, int, bool]]:
    """Danh sách (giá trị, số chữ số thập phân, có_hệ_số_triệu/nghìn) tìm thấy trong văn bản."""
    found: list[tuple[float, int, bool]] = []
    for m in _NUM_RE.finditer(str(text or "")):
        raw = m.group(0)
        neg = raw.lstrip().startswith(("-", "−", "–", "("))
        scale = 1.0
        sm = re.search(r"(triệu|trieu|million|tỷ|ty|billion|nghìn|nghin|ngàn|ngan|k)\s*$", raw, re.IGNORECASE)
        if sm:
            scale = _SCALE[sm.group(1).lower()]
        digits = re.search(r"\d[\d.,]*", raw)
        if not digits:
            continue
        tok = digits.group(0).rstrip(".,")
        try:
            for val, dec in _parse_token(tok):
                v = val * scale * (-1 if neg else 1)
                found.append((v, dec, scale != 1.0))
        except ValueError:
            continue
    return found


def number_matches(expected: float, found: list[tuple[float, int, bool]]) -> bool:
    """So khớp có chấp nhận sai số làm tròn; số âm được chấp nhận khi trình bày dạng 'lỗ 8,879.97'."""
    for val, dec, scaled in found:
        pairs = [(val, expected)]
        if expected < 0:
            pairs.append((abs(val), abs(expected)))
        for v, e in pairs:
            if e == 0:
                if abs(v) < 0.005:
                    return True
                continue
            rel = abs(v - e) / abs(e)
            if rel <= 1e-3:
                return True
            if scaled and rel <= 0.01:
                return True
            if not scaled and round(e, dec) == round(v, dec):
                return True
    return False


def split_items(cell: Any) -> list[list[str]]:
    """'a|b;c' -> [['a','b'],['c']]"""
    if cell is None or (isinstance(cell, float) and math.isnan(cell)) or str(cell).strip() == "":
        return []
    return [[alt.strip() for alt in item.split("|") if alt.strip()] for item in str(cell).split(";") if item.strip()]


# ----------------------------------------------------------------------------
# Chấm điểm một lượt chạy
# ----------------------------------------------------------------------------
def tools_used(result: dict) -> list[str]:
    return [str(t.get("tool", "")) for t in result.get("tool_results", [])]


def tool_has_error(result: dict) -> bool:
    for t in result.get("tool_results", []):
        r = t.get("result")
        if isinstance(r, dict) and (r.get("error") or r.get("status") == "error"):
            return True
    return False


def system_failure(result: dict) -> bool:
    return any(m in fold(result.get("answer", "")) for m in SYSTEM_FAILURE_MARKERS)


def grade(case: dict, result: dict) -> dict:
    """Trả về {'correct': bool, 'tool_ok': bool|None, 'detail': str, 'needs_review': bool}"""
    answer = str(result.get("answer", ""))
    check = str(case.get("Kiểu kiểm tra", "")).strip()
    used = tools_used(result)
    detail: list[str] = []

    # --- chọn đúng tool ---
    exp_tools = [t.strip() for t in str(case.get("Tool kỳ vọng") or "").split("|") if t.strip() and t.strip() != "nan"]
    tool_ok: Optional[bool] = None
    if exp_tools:
        tool_ok = any(t in used for t in exp_tools)
        if not tool_ok:
            detail.append(f"không gọi tool kỳ vọng {exp_tools} (đã gọi: {used or 'không'})")

    def numeric_ok() -> bool:
        items = split_items(case.get("Số liệu kỳ vọng"))
        found = extract_numbers(answer)
        ok_all = True
        for alts in items:
            hit = any(number_matches(float(a), found) for a in alts)
            if not hit:
                ok_all = False
                detail.append(f"thiếu số {'/'.join(alts)}")
        return ok_all

    def keyword_ok() -> bool:
        items = split_items(case.get("Từ khóa kỳ vọng"))
        ans = fold(answer)
        ok_all = True
        for alts in items:
            if not any(fold(a) in ans for a in alts):
                ok_all = False
                detail.append(f"thiếu từ khóa '{'/'.join(alts)}'")
        return ok_all

    correct = False
    needs_review = False
    if check == "numeric":
        correct = numeric_ok()
    elif check == "keyword":
        correct = keyword_ok()
    elif check == "numeric+keyword":
        n_ok, k_ok = numeric_ok(), keyword_ok()
        correct = n_ok and k_ok
    elif check == "chart":
        try:
            exp = json.loads(case.get("Tham số biểu đồ kỳ vọng") or "{}")
        except json.JSONDecodeError:
            exp = {}
        calls = [t for t in result.get("tool_results", []) if t.get("tool") == "generate_chart"]
        ok_call = False
        for t in calls:
            args, res = t.get("arguments", {}) or {}, t.get("result", {}) or {}
            if isinstance(res, dict) and not res.get("error") and all(str(args.get(k)) == str(v) for k, v in exp.items()):
                path = res.get("chart_path")
                if path and os.path.isfile(path):
                    ok_call = True
        correct = ok_call
        if not ok_call:
            detail.append(f"không có biểu đồ đúng tham số {exp}")
    elif check == "chart_error":
        calls = [t for t in result.get("tool_results", []) if t.get("tool") == "generate_chart"]
        got_error = any(isinstance(t.get("result"), dict) and t["result"].get("error") for t in calls)
        ans = fold(answer)
        admits = any(m in ans for m in ["khong the", "gia tri am", "khong phu hop", "bieu do cot", "xin loi", "loi"])
        fake_chart = any(isinstance(t.get("result"), dict) and t["result"].get("chart_path") for t in calls)
        correct = (not fake_chart) and admits
        needs_review = True
        if fake_chart:
            detail.append("tool vẫn tạo biểu đồ tròn dù có giá trị âm")
        if not got_error and not fake_chart:
            detail.append("agent không gọi generate_chart (kiểm tra tay)")
        if not admits:
            detail.append("câu trả lời không nêu rõ giới hạn/lỗi")
    elif check == "refuse":
        ans = fold(answer)
        marked = any(m in ans for m in REFUSE_MARKERS)
        correct = marked
        needs_review = True            # heuristic -> luôn nên duyệt tay
        if not marked:
            detail.append("không thấy dấu hiệu từ chối/nêu rõ không có dữ liệu")
    elif check == "empty":
        correct = "vui long" in fold(answer) and not result.get("tool_results")
        if not correct:
            detail.append("không trả thông báo yêu cầu nhập câu hỏi")
    else:
        detail.append(f"kiểu kiểm tra không hỗ trợ: {check}")

    if system_failure(result):
        correct = False
        detail.append("lỗi hệ thống: không gọi được LLM")
    return {"correct": bool(correct), "tool_ok": tool_ok, "detail": "; ".join(detail), "needs_review": needs_review}


# ----------------------------------------------------------------------------
# Chạy
# ----------------------------------------------------------------------------
def load_cases(path: str, ids: Optional[set], groups: Optional[set]) -> list[dict]:
    df = pd.read_excel(path, sheet_name="TestCases", dtype=object)
    df = df.where(pd.notna(df), None)
    cases = df.to_dict(orient="records")
    if ids:
        cases = [c for c in cases if c["ID"] in ids]
    if groups:
        cases = [c for c in cases if c["Nhóm"] in groups]
    return cases


def run_one(agent_fn: Callable[[str], dict], question: str) -> tuple[dict, float, Optional[str]]:
    t0 = time.perf_counter()
    exc = None
    try:
        result = agent_fn(question)
    except Exception as e:  # lỗi ngoài dự kiến của agent
        exc = f"{type(e).__name__}: {e}"
        result = {"success": False, "answer": f"[EXCEPTION] {exc}", "tool_results": [], "charts": [], "iterations": 0}
    return result, time.perf_counter() - t0, exc


def evaluate(cases: list[dict], agent_fn: Callable[[str], dict], runs: int, log=print,
             done: Optional[dict] = None, on_row: Optional[Callable[[dict], None]] = None) -> list[dict]:
    """done: {(ID, lần chạy): dòng kết quả} đã có từ lần chạy trước (bỏ qua, không chạy lại).
    on_row: được gọi ngay sau mỗi lượt để lưu checkpoint."""
    done = done or {}
    rows: list[dict] = []
    total = len(cases) * runs
    n = 0
    for case in cases:
        q = case["Câu hỏi"] or ""
        for r in range(1, runs + 1):
            n += 1
            if (case["ID"], r) in done:
                rows.append(done[(case["ID"], r)])
                log(f"[{n}/{total}] {case['ID']} lần {r}: bỏ qua (đã có kết quả)")
                continue
            result, dt, exc = run_one(agent_fn, q)
            g = grade(case, result)
            sys_err = bool(exc) or system_failure(result)
            tool_err = tool_has_error(result)
            row = {
                "ID": case["ID"], "Nhóm": case["Nhóm"], "Lần chạy": r, "Câu hỏi": q,
                "Kiểu kiểm tra": case["Kiểu kiểm tra"],
                "Kết quả": "Đúng" if g["correct"] else "Sai",
                "Đúng tool": ("" if g["tool_ok"] is None else ("Đúng" if g["tool_ok"] else "Sai")),
                "Thời gian (s)": round(dt, 3),
                "Số vòng lặp": result.get("iterations", 0),
                "Lỗi hệ thống": "Có" if sys_err else "Không",
                "Lỗi công cụ": "Có" if tool_err else "Không",
                "Tool đã gọi": ", ".join(tools_used(result)),
                "Lý do sai/ghi chú": g["detail"],
                "Cần duyệt tay": "Có" if g["needs_review"] else "",
                "Duyệt tay (Đúng/Sai)": "",
                "Câu trả lời": str(result.get("answer", ""))[:3000],
            }
            rows.append(row)
            if on_row:
                on_row(row)
            log(f"[{n}/{total}] {case['ID']} lần {r}: {row['Kết quả']} ({dt:.1f}s)"
                + (f" – {g['detail']}" if g["detail"] and not g["correct"] else ""))
    return rows


# ----------------------------------------------------------------------------
# Tổng hợp & xuất báo cáo
# ----------------------------------------------------------------------------
def pctl(values: list[float], p: float) -> float:
    if not values:
        return float("nan")
    s = sorted(values)
    k = (len(s) - 1) * p
    lo, hi = math.floor(k), math.ceil(k)
    return s[lo] + (s[hi] - s[lo]) * (k - lo)


def summarize(rows: list[dict]) -> dict:
    df = pd.DataFrame(rows)
    out: dict[str, Any] = {}
    out["total_runs"] = int(len(df))
    out["total_cases"] = int(df["ID"].nunique())
    out["accuracy"] = float((df["Kết quả"] == "Đúng").mean())
    t = df[df["Đúng tool"] != ""]
    out["tool_accuracy"] = float((t["Đúng tool"] == "Đúng").mean()) if len(t) else float("nan")
    lat = df["Thời gian (s)"].tolist()
    out["latency"] = {"mean": statistics.mean(lat), "median": statistics.median(lat),
                      "p95": pctl(lat, 0.95), "max": max(lat)}
    out["system_error_rate"] = float((df["Lỗi hệ thống"] == "Có").mean())
    out["tool_error_rate"] = float((df["Lỗi công cụ"] == "Có").mean())
    # nhất quán: mỗi câu, tỷ lệ lượt trùng kết quả đa số
    cons = []
    for _, g in df.groupby("ID"):
        if len(g) > 1:
            cons.append(g["Kết quả"].value_counts().iloc[0] / len(g))
    out["consistency"] = float(statistics.mean(cons)) if cons else float("nan")
    out["by_group"] = {
        grp: {"runs": int(len(g)), "accuracy": float((g["Kết quả"] == "Đúng").mean()),
              "mean_latency": float(g["Thời gian (s)"].mean())}
        for grp, g in df.groupby("Nhóm", sort=False)
    }
    return out


def write_report(rows: list[dict], meta: dict, path: str) -> None:
    """Xuất Excel: sheet Runs là dữ liệu thô; sheet TongHop dùng công thức tham chiếu sheet Runs."""
    cols = list(rows[0].keys())
    wb = Workbook()
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", start_color="1F4E78")
    base = Font(name="Arial", size=10)

    # --- Runs ---
    wr = wb.active
    wr.title = "Runs"
    for j, h in enumerate(cols, 1):
        c = wr.cell(row=1, column=j, value=h)
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(rows, 2):
        for j, h in enumerate(cols, 1):
            c = wr.cell(row=i, column=j, value=row[h] if row[h] != "" else None)
            c.font = base
            c.alignment = Alignment(vertical="top", wrap_text=(h in ("Câu hỏi", "Câu trả lời", "Lý do sai/ghi chú")))
    widths = {"ID": 8, "Nhóm": 13, "Lần chạy": 8, "Câu hỏi": 45, "Kiểu kiểm tra": 15, "Kết quả": 9, "Đúng tool": 9,
              "Thời gian (s)": 11, "Số vòng lặp": 9, "Lỗi hệ thống": 10, "Lỗi công cụ": 10, "Tool đã gọi": 28,
              "Lý do sai/ghi chú": 40, "Cần duyệt tay": 10, "Duyệt tay (Đúng/Sai)": 14, "Câu trả lời": 80}
    for j, h in enumerate(cols, 1):
        wr.column_dimensions[get_column_letter(j)].width = widths.get(h, 12)
    wr.freeze_panes = "E2"
    wr.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

    L = lambda name: get_column_letter(cols.index(name) + 1)
    last = len(rows) + 1
    rng = lambda name: f"Runs!${L(name)}$2:${L(name)}${last}"

    # --- TongHop ---
    ws = wb.create_sheet("TongHop", 0)
    ws["A1"] = "TỔNG HỢP ĐÁNH GIÁ AI AGENT"
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    info = [("Thời điểm chạy", meta["time"]), ("Model", meta["model"]), ("Số lần chạy mỗi câu", meta["runs"]),
            ("File test case", meta["cases_file"])]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=k).font = Font(name="Arial", size=10, bold=True)
        ws.cell(row=i, column=2, value=v).font = base
    r0 = 8
    ws.cell(row=r0 - 1, column=1, value="Chỉ số").font = head_font
    ws.cell(row=r0 - 1, column=2, value="Giá trị").font = head_font
    ws.cell(row=r0 - 1, column=3, value="Cách tính").font = head_font
    for c in range(1, 4):
        ws.cell(row=r0 - 1, column=c).fill = head_fill
    metrics = [
        ("Tổng số lượt chạy", f"=COUNTA({rng('ID')})", "0", "Số dòng trong sheet Runs"),
        ("Số lượt chạy đúng", f'=COUNTIF({rng("Kết quả")},"Đúng")', "0", "Kết quả = Đúng"),
        ("Độ chính xác", f"=B{r0 + 1}/B{r0}", "0.0%", "Lượt đúng / tổng lượt"),
        ("Chọn đúng công cụ", f'=IFERROR(COUNTIF({rng("Đúng tool")},"Đúng")/(COUNTIF({rng("Đúng tool")},"Đúng")+COUNTIF({rng("Đúng tool")},"Sai")),"n/a")', "0.0%", "Chỉ tính test có tool kỳ vọng"),
        ("Thời gian phản hồi TB (s)", f"=AVERAGE({rng('Thời gian (s)')})", "0.00", "Trung bình"),
        ("Thời gian phản hồi trung vị (s)", f"=MEDIAN({rng('Thời gian (s)')})", "0.00", "Trung vị"),
        ("Thời gian phản hồi P95 (s)", f"=PERCENTILE({rng('Thời gian (s)')},0.95)", "0.00", "Phân vị 95"),
        ("Thời gian phản hồi lớn nhất (s)", f"=MAX({rng('Thời gian (s)')})", "0.00", "Lớn nhất"),
        ("Tỷ lệ lỗi hệ thống", f'=COUNTIF({rng("Lỗi hệ thống")},"Có")/B{r0}', "0.0%", "Không gọi được LLM / exception"),
        ("Tỷ lệ lỗi công cụ", f'=COUNTIF({rng("Lỗi công cụ")},"Có")/B{r0}', "0.0%", "Tool trả về lỗi (kể cả lỗi có chủ đích ở test xử lý lỗi)"),
    ]
    for i, (name, f, fmt, note) in enumerate(metrics):
        r = r0 + i
        ws.cell(row=r, column=1, value=name).font = base
        c = ws.cell(row=r, column=2, value=f)
        c.font, c.number_format = base, fmt
        ws.cell(row=r, column=3, value=note).font = base
    r = r0 + len(metrics)
    ws.cell(row=r, column=1, value="Tính nhất quán (trung bình theo câu hỏi)").font = base
    c = ws.cell(row=r, column=2, value=f"=AVERAGE(CaHoi!$G$2:$G${len(meta['case_ids']) + 1})")
    c.font, c.number_format = base, "0.0%"
    ws.cell(row=r, column=3, value="Với mỗi câu: tỷ lệ lượt trùng kết quả đa số (cần runs ≥ 2)").font = base

    # theo nhóm
    gr = r + 3
    ws.cell(row=gr - 1, column=1, value="THEO NHÓM CÂU HỎI").font = Font(name="Arial", size=10, bold=True)
    for j, h in enumerate(["Nhóm", "Số lượt", "Số lượt đúng", "Độ chính xác", "Thời gian TB (s)"], 1):
        c = ws.cell(row=gr, column=j, value=h)
        c.font, c.fill = head_font, head_fill
    groups = list(dict.fromkeys(rw["Nhóm"] for rw in rows))
    for i, gname in enumerate(groups, 1):
        rr = gr + i
        ws.cell(row=rr, column=1, value=gname).font = base
        ws.cell(row=rr, column=2, value=f'=COUNTIF({rng("Nhóm")},A{rr})').font = base
        ws.cell(row=rr, column=3, value=f'=COUNTIFS({rng("Nhóm")},A{rr},{rng("Kết quả")},"Đúng")').font = base
        c = ws.cell(row=rr, column=4, value=f"=C{rr}/B{rr}")
        c.font, c.number_format = base, "0.0%"
        c = ws.cell(row=rr, column=5, value=f'=AVERAGEIF({rng("Nhóm")},A{rr},{rng("Thời gian (s)")})')
        c.font, c.number_format = base, "0.00"
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 58
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 16

    # --- CaHoi: tổng hợp theo từng câu ---
    wc = wb.create_sheet("CaHoi", 1)
    heads = ["ID", "Nhóm", "Câu hỏi", "Số lần chạy", "Số lần đúng", "Tỷ lệ đúng", "Nhất quán", "Thời gian TB (s)"]
    for j, h in enumerate(heads, 1):
        c = wc.cell(row=1, column=j, value=h)
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    qmap = {rw["ID"]: (rw["Nhóm"], rw["Câu hỏi"]) for rw in rows}
    for i, cid in enumerate(meta["case_ids"], 2):
        wc.cell(row=i, column=1, value=cid).font = base
        wc.cell(row=i, column=2, value=qmap[cid][0]).font = base
        c = wc.cell(row=i, column=3, value=qmap[cid][1] or "(rỗng)")
        c.font, c.alignment = base, Alignment(wrap_text=True, vertical="top")
        wc.cell(row=i, column=4, value=f'=COUNTIF({rng("ID")},A{i})').font = base
        wc.cell(row=i, column=5, value=f'=COUNTIFS({rng("ID")},A{i},{rng("Kết quả")},"Đúng")').font = base
        c = wc.cell(row=i, column=6, value=f"=E{i}/D{i}")
        c.font, c.number_format = base, "0%"
        # nhất quán = max(đúng, sai)/tổng
        c = wc.cell(row=i, column=7, value=f"=MAX(E{i},D{i}-E{i})/D{i}")
        c.font, c.number_format = base, "0%"
        c = wc.cell(row=i, column=8, value=f'=AVERAGEIF({rng("ID")},A{i},{rng("Thời gian (s)")})')
        c.font, c.number_format = base, "0.00"
    for col, w in zip("ABCDEFGH", [8, 14, 70, 11, 11, 11, 11, 15]):
        wc.column_dimensions[col].width = w
    wc.freeze_panes = "D2"

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


# ----------------------------------------------------------------------------
def real_agent_fn(model: str) -> Callable[[str], dict]:
    from src.agent import run_agent_loop  # import muộn để file này kiểm thử được khi chưa cài LangChain
    return lambda q: run_agent_loop(q, model=model)


def main() -> None:
    ap = argparse.ArgumentParser(description="Đánh giá AI Agent bằng bộ test case")
    ap.add_argument("--cases", default=DEFAULT_CASES)
    ap.add_argument("--runs", type=int, default=3, help="Số lần chạy lặp mỗi câu (đo tính nhất quán)")
    ap.add_argument("--ids", default="", help="Chỉ chạy các ID, ví dụ TC01,TC05")
    ap.add_argument("--groups", default="", help='Chỉ chạy các nhóm, ví dụ "Thống kê,Biểu đồ"')
    ap.add_argument("--model", default=None, help="Tên model Ollama (mặc định lấy từ src/agent.py)")
    ap.add_argument("--out", default=RESULT_DIR)
    ap.add_argument("--skip", default="", help="Bỏ qua các ID, ví dụ TC24,TC25,TC26,TC27")
    ap.add_argument("--fresh", action="store_true",
                    help="Xóa tiến độ cũ và chạy lại từ đầu (mặc định: chạy tiếp từ chỗ đã dừng)")
    args = ap.parse_args()

    if not os.path.isfile(args.cases):
        sys.exit(f"Không thấy {args.cases}. Chạy trước: python tests/build_test_cases.py")

    ids = {s.strip() for s in args.ids.split(",") if s.strip()} or None
    groups = {s.strip() for s in args.groups.split(",") if s.strip()} or None
    cases = load_cases(args.cases, ids, groups)
    skip = {s_.strip() for s_ in args.skip.split(",") if s_.strip()}
    cases = [c for c in cases if c["ID"] not in skip]
    if not cases:
        sys.exit("Không có test case nào khớp bộ lọc.")

    from src import agent as agent_module  # noqa: WPS433
    model = args.model or agent_module.MODEL_NAME
    agent_fn = real_agent_fn(model)

    # Checkpoint: mỗi lượt chạy xong được ghi ngay vào file .jsonl; chạy lại lệnh sẽ tiếp tục
    os.makedirs(args.out, exist_ok=True)
    ckpt = os.path.join(args.out, f"progress_{model.replace(':', '_').replace('/', '_')}_x{args.runs}.jsonl")
    if args.fresh and os.path.isfile(ckpt):
        os.remove(ckpt)
    done: dict = {}
    if os.path.isfile(ckpt):
        with open(ckpt, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line:
                    try:
                        r_ = json.loads(line)
                        done[(r_["ID"], r_["Lần chạy"])] = r_
                    except (json.JSONDecodeError, KeyError):
                        pass
        if done:
            print(f"Tìm thấy tiến độ cũ: {len(done)} lượt đã xong, sẽ chạy tiếp (dùng --fresh để chạy lại từ đầu).")

    def save_row(row: dict) -> None:
        with open(ckpt, "a", encoding="utf-8") as f:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")

    print(f"Chạy {len(cases)} test case x {args.runs} lần với model '{model}'...")
    rows = evaluate(cases, agent_fn, args.runs, done=done, on_row=save_row)
    summary = summarize(rows)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    meta = {"time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "model": model, "runs": args.runs,
            "cases_file": os.path.relpath(args.cases, BASE_DIR), "case_ids": [c["ID"] for c in cases]}
    xlsx = os.path.join(args.out, f"eval_{stamp}.xlsx")
    write_report(rows, meta, xlsx)
    with open(os.path.join(args.out, f"eval_{stamp}.json"), "w", encoding="utf-8") as f:
        json.dump({"meta": meta, "summary": summary, "runs": rows}, f, ensure_ascii=False, indent=2)

    lat = summary["latency"]
    print("\n===== KẾT QUẢ =====")
    print(f"Độ chính xác      : {summary['accuracy']:.1%}")
    print(f"Chọn đúng công cụ : {summary['tool_accuracy']:.1%}")
    print(f"Tính nhất quán    : {summary['consistency']:.1%}")
    print(f"Thời gian phản hồi: TB {lat['mean']:.1f}s | trung vị {lat['median']:.1f}s | P95 {lat['p95']:.1f}s | max {lat['max']:.1f}s")
    print(f"Lỗi hệ thống      : {summary['system_error_rate']:.1%} | Lỗi công cụ: {summary['tool_error_rate']:.1%}")
    for g, v in summary["by_group"].items():
        print(f"  - {g:<12} chính xác {v['accuracy']:.1%} ({v['runs']} lượt)")
    print(f"\nBáo cáo: {xlsx}")
    print("Lưu ý: các test 'refuse'/'chart_error' chấm bằng từ khóa -> cần duyệt tay cột 'Duyệt tay' trong sheet Runs.")


if __name__ == "__main__":
    main()
